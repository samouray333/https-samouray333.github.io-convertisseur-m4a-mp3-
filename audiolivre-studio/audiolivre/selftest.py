"""Autotest de bout en bout (utilisé par l'intégration continue et le support).

Lancement : ``AudioLivreStudio.exe --selftest rapport.json``
"""

from __future__ import annotations

import json
import logging
import math
import os
import platform
import subprocess
import sys
import tempfile
import time
import traceback
from pathlib import Path

import numpy as np

from . import __version__, paths

log = logging.getLogger(__name__)


def make_tone_engine():
    """Moteur interne de test (aucune dépendance) : produit une tonalité modulée."""
    from .core import audio
    from .core.engines.base import EngineInfo, ParamSpec, TTSEngine

    class ToneEngine(TTSEngine):
        info = EngineInfo(id="tone", name="Tonalité de test", tagline="Test", description="Test",
                          max_chars=200, params=[ParamSpec("speed", "Débit", 0.5, 2.0, 1.0)])

        def synthesize(self, text, voice, out_path, language, speed=1.0, seed=None):
            sr = 24000
            dur = max(0.5, len(text) / 14.5 / speed)
            t = np.arange(int(sr * dur)) / sr
            sig = 0.25 * np.sin(2 * math.pi * 180 * t) * (0.55 + 0.45 * np.sin(2 * math.pi * 3.3 * t))
            target = Path(out_path).with_suffix(".wav")
            audio.write_audio(target, sig.astype(np.float32), sr, subtype="PCM_16")
            return target

    return ToneEngine()


def run_selftest(report_path: Path) -> int:
    results: list[dict] = []
    critical_failed = False

    def check(name: str, critical: bool = True):
        def deco(fn):
            nonlocal critical_failed
            t0 = time.time()
            try:
                detail = fn()
                results.append({"check": name, "ok": True, "detail": detail, "seconds": round(time.time() - t0, 2)})
            except Exception as exc:  # noqa: BLE001
                results.append({"check": name, "ok": False, "critical": critical, "detail": str(exc),
                                "trace": traceback.format_exc()[-3000:], "seconds": round(time.time() - t0, 2)})
                if critical:
                    critical_failed = True
            return fn
        return deco

    work = Path(tempfile.mkdtemp(prefix="selftest-"))
    os.environ.setdefault("AUDIOLIVRE_HOME", str(work / "home"))

    @check("ffmpeg")
    def _():
        from .core import ffmpeg

        v = ffmpeg.version()
        assert "ffmpeg" in v.lower(), v
        for enc in ("libmp3lame", "aac", "flac", "libopus"):
            assert ffmpeg.has_encoder(enc), f"encodeur {enc} absent"
        return v

    @check("uv (installateur des moteurs)")
    def _():
        uv = paths.find_tool("uv")
        assert uv, "uv introuvable"
        from .core.ffmpeg import popen_kwargs

        out = subprocess.run([uv, "--version"], capture_output=True, text=True, **popen_kwargs())
        return out.stdout.strip()

    @check("script du processus moteur")
    def _():
        p = paths.worker_script()
        assert p.is_file(), f"absent : {p}"
        compile(p.read_text(encoding="utf-8"), str(p), "exec")
        return str(p)

    docs: dict[str, Path] = {}

    @check("import Word (.docx)")
    def _():
        import docx

        from .core import importers

        d = docx.Document()
        d.core_properties.title = "Livre de test"
        d.core_properties.author = "Autotest"
        d.add_heading("Chapitre 1 : Le départ", 1)
        d.add_paragraph("M. Dupont partit le 1er mai à 8h30 avec 12,50 € en poche.")
        d.add_heading("Chapitre 2 : Le retour", 1)
        d.add_paragraph("Il revint en 1984. — Bonjour ! dit-il.")
        p = work / "test.docx"
        d.save(str(p))
        docs["docx"] = p
        doc = importers.import_document(p)
        assert len(doc.chapters) == 2, [c.title for c in doc.chapters]
        return [c.title for c in doc.chapters]

    @check("import PDF")
    def _():
        import pymupdf as fitz

        from .core import importers

        pdf = fitz.open()
        for i in range(1, 4):
            page = pdf.new_page()
            page.insert_text((72, 90), f"Chapitre {i}", fontsize=22)
            page.insert_text((72, 140), f"Ceci est le texte du chapitre {i}. Il contient plusieurs phrases.",
                             fontsize=11)
            page.insert_text((72, 160), "Une deuxième ligne pour vérifier la reconstitution des paragraphes.",
                             fontsize=11)
        p = work / "test.pdf"
        pdf.save(str(p))
        doc = importers.import_document(p)
        assert len(doc.chapters) >= 3, [c.title for c in doc.chapters]
        return [c.title for c in doc.chapters]

    @check("normalisation du texte")
    def _():
        from .core.textproc import normalize_for_speech

        out = normalize_for_speech("M. Dupont a payé 12,50 € le 1er mai à 14h30 sous Louis XIV.")
        assert "Monsieur" in out and "douze euros cinquante" in out and "quatorze heures trente" in out, out
        assert "Louis quatorze" in out, out
        return out

    project_holder: dict = {}

    @check("production et assemblage")
    def _():
        from .core import engines, project_io, renderer
        from .core.importers import import_document
        from .core.models import Project, VoiceProfile
        from .core.voices import VoiceLibrary

        engines.register_engine(make_tone_engine())
        lib = VoiceLibrary(work / "voices")
        v = lib.save(VoiceProfile(name="Test", engine="tone"))
        doc = import_document(docs["docx"])
        pr = Project(metadata=doc.metadata, chapters=doc.chapters, narrator_voice_id=v.id)
        project_io.save_project(pr, work / "projet" / "Test.alsproj")
        plan = renderer.build_plan(pr, lib)
        rep = renderer.Renderer(pr, lib).render(plan)
        assert not rep.failed_segments, rep.failed_segments
        assert all(renderer.chapter_is_current(pr, rc) for rc in plan)
        project_holder.update(pr=pr, plan=plan)
        return f"{len(plan)} pistes, {sum(len(rc.segments) for rc in plan)} passages"

    @check("mastering ACX et export M4B/MP3")
    def _():
        from .core import ffmpeg, renderer
        from .core.exporter import ExportItem, export_audiobook
        from .core.mastering import MasteringOptions

        pr, plan = project_holder["pr"], project_holder["plan"]
        pr.export.formats = ["m4b", "mp3_chapters", "flac"]
        items = [ExportItem(rc.title, renderer.chapter_output(pr, rc.id), rc.kind) for rc in plan]
        res = export_audiobook(items, pr.metadata, pr.export, MasteringOptions(), work / "export")
        m4b = next(f for f in res.files if f.suffix == ".m4b")
        info = ffmpeg.probe(m4b)
        assert len(info.get("chapters", [])) == len(plan), info.get("chapters")
        assert res.acx_ok, [(t, s.acx_checks()) for t, s in res.stats]
        return [f.name for f in res.files]

    @check("voix Windows (SAPI)", critical=False)
    def _():
        from .core.engines.sapi import SapiEngine
        from .core.models import VoiceProfile

        e = SapiEngine()
        st, msg = e.status()
        if st != "ready":
            return f"ignoré : {msg}"
        voices = e.list_builtin_voices()
        if not voices:
            return "aucune voix installée"
        out = e.synthesize("Bonjour, ceci est un test.", VoiceProfile(name="t", engine="sapi",
                                                                       engine_voice=voices[0].id),
                           work / "sapi.wav", "fr")
        assert out.exists() and out.stat().st_size > 1000
        return f"{len(voices)} voix ; {voices[0].name}"

    @check("voix Microsoft en ligne (Edge)", critical=False)
    def _():
        from .core.engines.edge import EdgeEngine
        from .core.models import VoiceProfile

        e = EdgeEngine()
        out = e.synthesize("Bonjour, ceci est un test des voix neuronales.",
                           VoiceProfile(name="Denise", engine="edge", engine_voice="fr-FR-DeniseNeural"),
                           work / "edge.mp3", "fr")
        from .core import audio

        data, sr = audio.read_audio(out)
        assert len(data) / sr > 1.0
        return f"{len(data) / sr:.1f} s"

    @check("interface graphique")
    def _():
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        from PySide6.QtWidgets import QApplication

        app = QApplication.instance() or QApplication(sys.argv[:1])
        from .ui import theme

        theme.apply(app, "dark")
        from .ui.main_window import MainWindow

        win = MainWindow()
        win.ctx.set_project(project_holder.get("pr") or win.ctx.project)
        for key in win.pages:
            win.show_page(key)
            app.processEvents()
        win.ctx.dirty = False
        win.close()
        return f"{len(win.pages)} pages"

    report = {
        "version": __version__, "python": sys.version, "platform": platform.platform(),
        "frozen": paths.is_frozen(), "ok": not critical_failed, "results": results,
    }
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    for r in results:
        log.info("[%s] %s — %s", "OK" if r["ok"] else "ÉCHEC", r["check"], r["detail"])
    return 0 if not critical_failed else 1
