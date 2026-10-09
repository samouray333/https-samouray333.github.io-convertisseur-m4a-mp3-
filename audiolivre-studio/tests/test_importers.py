import pytest

from audiolivre.core import importers


def test_docx_headings(tmp_path):
    docx = pytest.importorskip("docx")
    d = docx.Document()
    d.core_properties.title = "Mon livre"
    d.core_properties.author = "Jeanne"
    d.add_paragraph("Préface courte.")
    d.add_heading("Chapitre 1", 1)
    d.add_paragraph("Premier texte.")
    d.add_heading("Une section", 2)
    d.add_paragraph("Suite.")
    d.add_heading("Chapitre 2", 1)
    d.add_paragraph("Second texte.")
    p = tmp_path / "livre.docx"
    d.save(str(p))
    doc = importers.import_document(p)
    assert doc.metadata.title == "Mon livre"
    assert doc.metadata.author == "Jeanne"
    assert [c.title for c in doc.chapters] == ["Chapitre 1", "Chapitre 2"]
    assert "Préface courte." in doc.chapters[0].text
    assert "# Une section" in doc.chapters[0].text


def test_pdf_chapters_and_headers(tmp_path):
    fitz = pytest.importorskip("pymupdf")
    pdf = fitz.open()
    for i in range(1, 5):
        page = pdf.new_page()
        page.insert_text((72, 30), "Mon Livre — en-tête", fontsize=9)
        page.insert_text((72, 100), f"Chapitre {i}", fontsize=22)
        page.insert_text((72, 150), f"Texte du chapitre {i}, avec une phrase qui conti-", fontsize=11)
        page.insert_text((72, 165), "nue sur la ligne suivante.", fontsize=11)
        page.insert_text((300, 820), str(i), fontsize=9)
    p = tmp_path / "livre.pdf"
    pdf.save(str(p))
    doc = importers.import_document(p)
    assert [c.title for c in doc.chapters] == [f"Chapitre {i}" for i in range(1, 5)]
    assert "continue sur la ligne suivante" in doc.chapters[0].text
    assert "en-tête" not in doc.chapters[1].text


def test_text_chapter_patterns(tmp_path):
    p = tmp_path / "livre.txt"
    p.write_text("Introduction du livre avec suffisamment de mots pour former une ouverture digne de ce nom, "
                 "car il faut plus de vingt-cinq mots pour cela, voilà qui est fait maintenant.\n\n"
                 "CHAPITRE I\n\nLe début.\n\nChapitre II\n\nLa suite.\n", encoding="utf-8")
    doc = importers.import_document(p)
    titles = [c.title for c in doc.chapters]
    assert titles[0] == "Ouverture"
    assert "CHAPITRE I" in titles and "Chapitre II" in titles


def test_markdown(tmp_path):
    p = tmp_path / "livre.md"
    p.write_text("# Partie 1\n\nDu **texte** avec un [lien](http://x).\n\n# Partie 2\n\nFin.\n", encoding="utf-8")
    doc = importers.import_document(p)
    assert [c.title for c in doc.chapters] == ["Partie 1", "Partie 2"]
    assert "Du texte avec un lien." in doc.chapters[0].text


def test_html(tmp_path):
    p = tmp_path / "page.html"
    p.write_text("<html><body><h1>Titre A</h1><p>Un.</p><h1>Titre B</h1><p>Deux.</p></body></html>",
                 encoding="utf-8")
    doc = importers.import_document(p)
    assert [c.title for c in doc.chapters] == ["Titre A", "Titre B"]


def test_epub(tmp_path):
    epub = pytest.importorskip("ebooklib.epub")
    book = epub.EpubBook()
    book.set_identifier("id1")
    book.set_title("Roman")
    book.set_language("fr")
    book.add_author("Victor")
    chs = []
    for i in range(1, 3):
        c = epub.EpubHtml(title=f"Ch {i}", file_name=f"c{i}.xhtml", lang="fr")
        c.content = f"<h1>Chapitre {i}</h1><p>Texte numéro {i} du roman, assez long pour compter.</p>"
        book.add_item(c)
        chs.append(c)
    book.toc = chs
    book.spine = chs
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    p = tmp_path / "roman.epub"
    epub.write_epub(str(p), book)
    doc = importers.import_document(p)
    assert doc.metadata.title == "Roman"
    assert [c.title for c in doc.chapters][-2:] == ["Chapitre 1", "Chapitre 2"]


def test_unsupported(tmp_path):
    p = tmp_path / "x.doc"
    p.write_bytes(b"x")
    with pytest.raises(importers.ImportErrorUser):
        importers.import_document(p)


def test_split_by_size_without_titles(tmp_path):
    p = tmp_path / "long.txt"
    p.write_text("\n\n".join(["Un paragraphe de vingt mots " * 4] * 600), encoding="utf-8")
    doc = importers.import_document(p)
    assert len(doc.chapters) >= 2
    assert doc.chapters[0].title == "Partie 1"


def test_docx_dash_bullets_become_dialogues(tmp_path):
    docx = pytest.importorskip("docx")
    from docx.oxml import parse_xml

    d = docx.Document()
    w = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    numbering = d.part.numbering_part.element
    numbering.append(parse_xml(f'<w:abstractNum {w} w:abstractNumId="90"><w:lvl w:ilvl="0"><w:numFmt w:val="bullet"/>'
                               f'<w:lvlText w:val="-"/></w:lvl></w:abstractNum>'))
    numbering.append(parse_xml(f'<w:num {w} w:numId="90"><w:abstractNumId w:val="90"/></w:num>'))
    d.add_heading("Chapitre 1", 1)
    d.add_paragraph("Marie entra.")
    for line in ("Bonjour, dit-elle.", "Salut !"):
        para = d.add_paragraph(line)
        para._p.get_or_add_pPr().append(parse_xml(f'<w:numPr {w}><w:ilvl w:val="0"/><w:numId w:val="90"/></w:numPr>'))
    d.add_paragraph("Liste ordinaire", style="List Bullet")
    p = tmp_path / "dialogues.docx"
    d.save(str(p))
    text = importers.import_document(p).chapters[0].text
    assert "— Bonjour, dit-elle." in text and "— Salut !" in text
    assert "— Liste ordinaire" not in text and "— Marie entra." not in text
